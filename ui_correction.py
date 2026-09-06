"""
Streamlit UI for the "Regenerate from existing report" workflow.

Extracts every piece of information out of a finished report (title,
date & time, venue, description, poster and event photos), lets the user
fix whatever is wrong, and regenerates a brand-new, correctly aligned
document from the template.

Kept as a separate module so it can be added/removed without affecting
the template-based generation flow in app.py.
"""

import streamlit as st
from pathlib import Path
import tempfile
from io import BytesIO

from docx import Document

from engine.template_loader import load_template
from engine.content_inserter import replace_placeholders
from engine.extractor import (
    extract_report_data,
    build_generation_data,
    rewrite_dates_in_text,
)
from engine.preview import render_docx_preview, show_preview

from pathlib import Path as _Path

_BASE = _Path(__file__).parent
TEMPLATE = _BASE / "templates" / "final_sample_report.docx"


def render_correction_ui():

    st.subheader("Regenerate from Existing Report")

    st.write(
        "Upload a finished report. The tool pulls every piece of "
        "information out of it (title, date & time, venue, "
        "description, poster and event photos). Fix whatever is "
        "wrong below, then regenerate — you get a brand-new document "
        "with proper alignment."
    )

    uploaded = st.file_uploader(
        "Finished report (.docx)",
        type=["docx"],
    )

    extract = st.button(
        "Extract data from report",
        use_container_width=True,
    )

    if uploaded and extract:

        try:

            blob = uploaded.getvalue()

            document = Document(BytesIO(blob))

            extracted = extract_report_data(document)

            st.session_state["re_extracted"] = {
                "date_time": extracted["date_time"] or "",
                "venue": extracted["venue"] or "",
                "title": extracted["title"] or "",
                "description": extracted["description"] or "",
                "poster_image": extracted["poster_image"],
                "photo1_image": extracted["photo1_image"],
                "photo2_image": extracted["photo2_image"],
            }

            st.session_state["re_original_name"] = uploaded.name

            st.success("Report data extracted!")

        except Exception as e:

            st.exception(e)

    extracted = st.session_state.get("re_extracted")

    if extracted is None:

        st.info(
            "Upload a finished report and press "
            "'Extract data from report' to begin."
        )
        return

    original_name = st.session_state.get("re_original_name", "report.docx")

    # ------------------------------------------------------------
    # TEXT FIELDS (prefilled with the extracted values)
    # ------------------------------------------------------------

    st.subheader("Report Details")

    st.caption(
        "These values were pulled out of your report. "
        "Edit any of them: Title, Date & Time, Venue and "
        "Description will be placed into the regenerated "
        "document."
    )

    data = {}

    for field, label, placeholder in (
        ("title", "Event Title", "AUTOMATION PROGRAMMING WORKSHOP"),
        ("venue", "Venue", "CSE LAB – 1"),
    ):

        data[field] = st.text_input(
            label,
            value=extracted[field],
            placeholder=placeholder,
        )

    data["date_time"] = st.text_input(
        "Date & Time",
        value=extracted["date_time"],
        placeholder="30-01-2026 & 01.30 am to 5 pm",
    )

    data["description"] = st.text_area(
        "Event Description",
        value=extracted["description"],
        height=300,
    )

    st.divider()

    # ------------------------------------------------------------
    # IMAGES (extracted images are reused unless replaced)
    # ------------------------------------------------------------

    st.subheader("Event Images")

    st.caption(
        "The poster and photos were extracted from the report. "
        "Upload a replacement to use it instead; leave empty to "
        "reuse the extracted image."
    )

    replace_poster = st.file_uploader(
        "Replace Poster (used for medium + full poster)",
        type=["jpg", "jpeg", "png"],
    )

    image_overrides = {}

    if replace_poster is not None:
        image_overrides["poster"] = BytesIO(replace_poster.getbuffer())

    cols = st.columns(2)

    with cols[0]:

        st.write("**Extracted poster**")

        if extracted["poster_image"]:
            st.image(BytesIO(extracted["poster_image"]), width=260)
        else:
            st.write("_No poster found._")

    with cols[1]:

        replace_photo1 = st.file_uploader(
            "Replace Event Photo 1",
            type=["jpg", "jpeg", "png"],
        )

        if replace_photo1 is not None:
            image_overrides["photo1"] = BytesIO(
                replace_photo1.getbuffer()
            )

    cols = st.columns(2)

    with cols[0]:

        st.write("**Extracted photo 1**")

        if extracted["photo1_image"]:
            st.image(BytesIO(extracted["photo1_image"]), width=260)
        else:
            st.write("_No photo found._")

    with cols[1]:

        replace_photo2 = st.file_uploader(
            "Replace Event Photo 2",
            type=["jpg", "jpeg", "png"],
        )

        if replace_photo2 is not None:
            image_overrides["photo2"] = BytesIO(
                replace_photo2.getbuffer()
            )

    cols = st.columns(2)

    with cols[0]:

        st.write("**Extracted photo 2**")

        if extracted["photo2_image"]:
            st.image(BytesIO(extracted["photo2_image"]), width=260)
        else:
            st.write("_No photo found._")

    # ------------------------------------------------------------
    # GENERATE
    # ------------------------------------------------------------

    st.divider()

    generate = st.button(
        "🚀 Regenerate Report",
        type="primary",
        use_container_width=True,
    )

    if not generate:
        return

    # ------------------------------------------------------------
    # VALIDATION
    # ------------------------------------------------------------

    missing = []

    if not data["title"].strip():
        missing.append("Event Title")
    if not data["date_time"].strip():
        missing.append("Date & Time")
    if not data["venue"].strip():
        missing.append("Venue")
    if not data["description"].strip():
        missing.append("Event Description")

    if missing:

        st.error("Please provide: " + ", ".join(missing))
        return

    try:

        with st.spinner("Regenerating report..."):

            description = rewrite_dates_in_text(
                data["description"],
                extracted["date_time"],
                data["date_time"],
            )

            if description != data["description"]:
                st.caption(
                    "The date mentions inside the description were "
                    "rewritten to match the new Date & Time."
                )

            generation_data = build_generation_data(
                {
                    "date_time": data["date_time"],
                    "venue": data["venue"],
                    "title": data["title"],
                    "description": description,
                    "poster_image": extracted["poster_image"],
                    "photo1_image": extracted["photo1_image"],
                    "photo2_image": extracted["photo2_image"],
                },
                image_overrides,
            )

            document = load_template(TEMPLATE)

            document = replace_placeholders(
                document,
                generation_data,
            )

            with tempfile.TemporaryDirectory() as temp_dir:

                temp_dir = Path(temp_dir)

                file_stem = Path(original_name).stem
                output_path = temp_dir / f"regenerated_{file_stem}.docx"

                document.save(output_path)

                output_bytes = output_path.read_bytes()

                pages = render_docx_preview(output_bytes, temp_dir)

        st.success("✅ Report regenerated successfully!")

        st.download_button(
            label="⬇️ Download Regenerated Report",
            data=output_bytes,
            file_name=f"regenerated_{file_stem}.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            use_container_width=True,
        )

        show_preview(pages)

    except Exception as e:

        st.error(
            "Something went wrong while regenerating "
            "the report."
        )

        st.exception(e)