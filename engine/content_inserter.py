from copy import deepcopy
from pathlib import Path

from docx.text.paragraph import Paragraph
from docx.shared import Inches
from PIL import Image


IMAGE_TYPES = {
    "{{POSTER_MEDIUM}}": "POSTER_MEDIUM",
    "{{EVENT_PHOTO_1}}": "EVENT_PHOTO_1",
    "{{EVENT_PHOTO_2}}": "EVENT_PHOTO_2",
    "{{POSTER_FULL}}": "POSTER_FULL",
}


def replace_placeholders(document, data):
    """
    Final report template engine.

    Handles:
    - Normal text placeholders
    - Multi-paragraph content
    - Line breaks
    - Images
    - Automatic section page breaks
    - Long event titles
    - Poster / Photos / Full Poster sections
    """

    # Process normal document paragraphs
    for paragraph in list(document.paragraphs):
        _process_paragraph(paragraph, data)

    # Process placeholders inside tables
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in list(cell.paragraphs):
                    _process_paragraph(paragraph, data)

    return document


# ============================================================
# PARAGRAPH PROCESSING
# ============================================================

def _process_paragraph(paragraph, data):

    full_text = "".join(run.text for run in paragraph.runs).strip()

    if not full_text:
        return

    # --------------------------------------------------------
    # IMAGE PLACEHOLDERS
    # --------------------------------------------------------

    for placeholder, image_type in IMAGE_TYPES.items():

        if placeholder in full_text:

            image_path = data.get(image_type)

            if image_path:
                _insert_image(
                    paragraph,
                    image_path,
                    image_type
                )

            return

    # --------------------------------------------------------
    # NORMAL TEXT
    # --------------------------------------------------------

    _replace_text_placeholder(paragraph, data)


# ============================================================
# TEXT PLACEHOLDER REPLACEMENT
# ============================================================

def _replace_text_placeholder(paragraph, data):

    image_keys = {
        "POSTER_MEDIUM",
        "EVENT_PHOTO_1",
        "EVENT_PHOTO_2",
        "POSTER_FULL",
    }

    for key, value in data.items():

        if key in image_keys:
            continue

        placeholder = "{{" + key + "}}"

        full_text = "".join(
            run.text for run in paragraph.runs
        )

        if placeholder not in full_text:
            continue

        _replace_text(
            paragraph,
            placeholder,
            str(value)
        )


# ============================================================
# TEXT WITH MULTIPLE PARAGRAPHS
# ============================================================

def _replace_text(paragraph, placeholder, value):

    # Normalize literal backslash-n sequences (e.g. "\n" typed
    # as two characters) into real newlines so they are rendered
    # as line breaks / paragraph splits instead of visible text.
    value = value.replace("\\n", "\n")

    # Replace only the requested placeholder
    full_text = "".join(
        run.text for run in paragraph.runs
    )

    value = full_text.replace(
        placeholder,
        value
    )

    parts = value.split("\n\n")

    _set_paragraph_content(
        paragraph,
        parts[0]
    )

    current_paragraph = paragraph

    for part in parts[1:]:

        new_xml = deepcopy(paragraph._p)

        # Remove existing runs
        for child in list(new_xml):
            if child.tag.endswith("}r"):
                new_xml.remove(child)

        current_paragraph._p.addnext(new_xml)

        new_paragraph = Paragraph(
            new_xml,
            current_paragraph._parent
        )

        _set_paragraph_content(
            new_paragraph,
            part
        )

        current_paragraph = new_paragraph


# ============================================================
# TEXT + SINGLE LINE BREAKS
# ============================================================

def _set_paragraph_content(paragraph, text):

    # Clear existing text
    for run in paragraph.runs:
        run.text = ""

    if paragraph.runs:
        run = paragraph.runs[0]
    else:
        run = paragraph.add_run()

    lines = text.split("\n")

    for index, line in enumerate(lines):

        run.add_text(line)

        if index < len(lines) - 1:
            run.add_break()


# ============================================================
# IMAGE INSERTION
# ============================================================

def _insert_image(paragraph, image, image_type):

    is_file_like = hasattr(image, "read") and hasattr(image, "seek")

    if is_file_like:
        image_path = None
        image.seek(0)
    else:
        image_path = Path(image)
        if not image_path.exists():
            raise FileNotFoundError(
                f"Image not found: {image_path}"
            )

    # Remove placeholder
    for run in paragraph.runs:
        run.text = ""

    # Images are always centered
    paragraph.alignment = 1

    # --------------------------------------------------------
    # SECTION PAGE BREAKS
    # --------------------------------------------------------

    if image_type == "POSTER_MEDIUM":

        # Start Poster section on a fresh page.
        # Transfer the break to the "Poster" heading (if present),
        # so the heading is not left stranded on the previous page.
        if not _break_previous_heading(
            paragraph,
            "Poster"
        ):
            paragraph.paragraph_format.page_break_before = True

    elif image_type == "EVENT_PHOTO_1":

        # Start Photos section on a fresh page.
        # Transfer the break to the "Photos" heading (if present).
        if not _break_previous_heading(
            paragraph,
            "Photos"
        ):
            paragraph.paragraph_format.page_break_before = True

    elif image_type == "POSTER_FULL":

        # Full poster always starts on a fresh page.
        paragraph.paragraph_format.page_break_before = True

    # EVENT_PHOTO_2 intentionally gets NO page break.
    # Therefore both photos stay on the same page.

    # --------------------------------------------------------
    # IMAGE SIZE
    # --------------------------------------------------------

    if image_type == "POSTER_MEDIUM":

        max_width = Inches(5.7)
        max_height = Inches(5.2)

    elif image_type == "POSTER_FULL":

        max_width = Inches(6.3)
        max_height = Inches(8.8)

    elif image_type in {
        "EVENT_PHOTO_1",
        "EVENT_PHOTO_2",
    }:

        max_width = Inches(4.8)
        max_height = Inches(3.3)

    else:
        return

    # --------------------------------------------------------
    # READ IMAGE DIMENSIONS
    # --------------------------------------------------------

    if is_file_like:
        image.seek(0)

    with Image.open(image) as img:

        width_px, height_px = img.size

    aspect_ratio = width_px / height_px

    # Fit by width first
    width_inches = max_width.inches
    height_inches = width_inches / aspect_ratio

    # If too tall, fit by height
    if height_inches > max_height.inches:

        height_inches = max_height.inches
        width_inches = height_inches * aspect_ratio

    # --------------------------------------------------------
    # INSERT
    # --------------------------------------------------------

    run = (
        paragraph.runs[0]
        if paragraph.runs
        else paragraph.add_run()
    )

    if is_file_like:
        image.seek(0)
        run.add_picture(
            image,
            width=Inches(width_inches),
            height=Inches(height_inches)
        )
    else:
        run.add_picture(
            str(image_path),
            width=Inches(width_inches),
            height=Inches(height_inches)
        )


# ============================================================
# MOVE SECTION HEADING TO NEW PAGE
# ============================================================

def _break_previous_heading(paragraph, heading_text):

    """
    Finds the immediately preceding paragraph.

    If it contains the expected section heading,
    the page break is applied to that heading instead
    of leaving the heading stranded on the previous page.

    Returns True when the break was transferred to the heading,
    False otherwise.
    """

    previous = paragraph._p.getprevious()

    if previous is None:
        return False

    try:
        previous_paragraph = Paragraph(
            previous,
            paragraph._parent
        )
    except Exception:
        return False

    previous_text = (
        "".join(
            run.text
            for run in previous_paragraph.runs
        )
        .strip()
        .rstrip(":")
        .strip()
    )

    heading_text = heading_text.strip().rstrip(":").strip()

    if previous_text.lower() == heading_text.lower():

        previous_paragraph.paragraph_format.page_break_before = True

        return True

    return False